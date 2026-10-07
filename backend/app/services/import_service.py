# -*- coding: utf-8 -*-
"""
Importação de extrato e conciliação (Fatia 3 — ver docs/finance/spec.md).

Regras decididas na sabatina de 21-22/09:

- **Duplicata não entra.** O que é suspeito de duplicidade espera numa antessala
  (a própria tabela de transações, com `ingest_state`), e nada que espera
  aparece em saldo, relatório ou listagem padrão.
- **Entrada parcial**: as linhas limpas entram na hora; só as suspeitas esperam.
- **Conciliar funde**: o registro que fica mantém `id`, categoria e edições, mas
  absorve os identificadores da linha importada. Sem isso, a mesma linha
  voltaria como suspeita a cada reimportação do período.
- **Reimportar reconhece** o que já entrou e não traz de novo.
- Linha importada é fato consumado: nasce com `settled_at` preenchido, porque o
  banco só informa o que já aconteceu.

Decisões de implementação que o spec deixava em aberto, e por quê:

1. **O hash leva um contador de ocorrência.** Dois cafés de R$ 8 no mesmo dia
   geram a mesma base (descrição + valor + data + conta). Sem o contador, a
   segunda linha seria descartada como "já importada". Com ele, reimportar
   continua idempotente e as linhas legítimas repetidas ficam distintas.
2. **Só é candidato a duplicata quem NÃO tem FITID.** Registro digitado à mão,
   gerado por série ou importado de CSV. Duas compras iguais vindas do banco,
   cada uma com seu FITID, são transações distintas — o banco já disse — e
   suspeitar entre elas encheria a fila de falso positivo. Isso inclui o
   registro **previsto** de uma série: é assim que "previsto vira realizado"
   acontece na conciliação.
3. **Palavra genérica não casa** ("compra", "pix", "pagamento"…). Casaria quase
   tudo com o mesmo valor e a mesma data.
4. **Cada candidato é reivindicado por uma linha só.** Duas linhas iguais no
   extrato e um lançamento manual: uma é suspeita, a outra entra limpa.
5. **A conta é escolhida pelo usuário e entra no hash e no id externo** — não o
   ACCTID do arquivo. O CSV não tem conta; com a conta escolhida, o mesmo extrato
   em CSV e em OFX gera o MESMO hash e é reconhecido como já importado. E duas
   contas com a mesma linha (uma transferência aparece nas duas pontas) continuam
   distintas, sem nenhum campo tirado do hash.
6. **O ACCTID do arquivo serve de aviso, não de identidade.** Guardado na conta na
   primeira importação; se depois chegar um arquivo com outro, a prévia avisa
   "parece ser de outra conta" — o erro de mandar o extrato do Bradesco para a
   conta do C6 não pode passar em silêncio.
7. **Erro de falso negativo custa mais que o de falso positivo.** Suspeitar de
   mais vira uma linha a conferir; suspeitar de menos duplica dinheiro em silêncio.
"""

import csv
import hashlib
import io
import re
import unicodedata
from dataclasses import dataclass
from datetime import date, datetime, timezone
from typing import Any, Dict, List, Optional, Tuple

from app.services.data_service import DataService, em_transacao

# Distância máxima entre a data do extrato e a do lançamento existente.
JANELA_DIAS = 3

# Palavras que aparecem em quase toda descrição bancária e por isso não
# identificam nada. Ficam de fora do casamento por palavra-chave.
STOPWORDS = frozenset(
    {
        "compra", "compras", "cartao", "credito", "debito", "pagamento", "pagto",
        "pgto", "transferencia", "transf", "pix", "boleto", "saque", "tarifa",
        "recebimento", "deposito", "enviado", "enviada", "recebido", "recebida",
        "referente", "docto", "documento", "banco", "conta", "cta",
    }
)

FORMATOS_ACEITOS = ("ofx", "qfx", "csv", "txt", "xlsx")


class ArquivoInvalido(ValueError):
    """Arquivo que não dá pra importar — a mensagem é mostrada ao usuário."""


@dataclass(frozen=True, eq=False)
class Destino:
    """
    Para onde o extrato vai: uma Conta OU um Cartão (issue #27).

    Concentra tudo o que muda entre os dois, para o resto da importação não
    perguntar "é cartão?" em cada passo: o prefixo do hash, a coleção onde
    mora o identificador do banco, quais lançamentos existentes podem ser
    duplicata, o texto do aviso e o formato do resumo na resposta.
    """

    eh_cartao: bool
    registro: Dict[str, Any]

    @classmethod
    def resolver(cls, user_id: int, account_id: Optional[int], card_id: Optional[int]) -> "Destino":
        """Exatamente um dos dois ids precisa vir, e o registro precisa ser do usuário."""
        if (account_id is None) == (card_id is None):
            raise ArquivoInvalido("Escolha o destino da importação: uma conta ou um cartão.")
        eh_cartao = card_id is not None
        alvo_id = card_id if eh_cartao else account_id
        registro = next(
            (
                r
                for r in DataService.load_json("cards" if eh_cartao else "accounts")
                if r["id"] == alvo_id and r["user_id"] == user_id
            ),
            None,
        )
        if registro is None:
            raise LookupError("Cartão não encontrado." if eh_cartao else "Conta não encontrada.")
        return cls(eh_cartao, registro)

    @property
    def id(self) -> int:
        return self.registro["id"]

    @property
    def tipo(self) -> str:
        """Também é a chave do resumo na resposta da API: `conta` ou `cartao`."""
        return "cartao" if self.eh_cartao else "conta"

    @property
    def colecao(self) -> str:
        return "cards" if self.eh_cartao else "accounts"

    @property
    def account_id(self) -> Optional[int]:
        return None if self.eh_cartao else self.id

    @property
    def card_id(self) -> Optional[int]:
        return self.id if self.eh_cartao else None

    @property
    def prefixo_hash(self) -> str:
        """Separa o hash de conta do de cartão: o mesmo texto nos dois são lançamentos distintos."""
        return ("card%d" if self.eh_cartao else "acc%d") % self.id

    def resumo(self) -> Dict[str, Any]:
        resumo = {"id": self.id, "name": self.registro["name"]}
        if not self.eh_cartao:
            resumo["kind"] = self.registro["kind"]
        return resumo

    def pode_ser_duplicata(self, t: Dict[str, Any]) -> bool:
        """
        Lançamento existente que pode ser a mesma compra de uma linha do extrato.

        No Cartão, a compra manual do MESMO cartão. Na Conta, a da mesma conta
        — ou de nenhuma: o lançamento digitado à mão ainda não pergunta a conta,
        e um candidato de OUTRA conta não pode ser fundido com esta linha.
        Compra de cartão nunca serve para a Conta: não tem conta.
        """
        if self.eh_cartao:
            return t.get("card_id") == self.id
        return not t.get("card_id") and t.get("account_id") in (None, self.id)

    def aviso_de_arquivo_trocado(self, acctid: Optional[str]) -> Optional[str]:
        """O identificador do banco no arquivo difere do já guardado: provável arquivo de outro destino."""
        guardado = self.registro.get("file_ref")
        if not acctid or not guardado or acctid == guardado:
            return None
        outro, certo = ("outro cartão", "o cartão certo") if self.eh_cartao else ("outra conta", "a conta certa")
        return (
            "Este arquivo parece ser de %s (identificador %s; \"%s\" já recebeu arquivos do %s). "
            "Confira se escolheu %s." % (outro, acctid, self.registro["name"], guardado, certo)
        )


# --------------------------------------------------------------------------- #
# Normalização
# --------------------------------------------------------------------------- #


def normalizar(texto: Optional[str]) -> str:
    """Minúsculas, sem acento, espaços colapsados — a base de hash e casamento."""
    if not texto:
        return ""
    sem_acento = unicodedata.normalize("NFKD", texto)
    sem_acento = "".join(c for c in sem_acento if not unicodedata.combining(c))
    return re.sub(r"\s+", " ", sem_acento.lower()).strip()


def _centavos(valor: float) -> int:
    """Valor em centavos inteiros: comparar float por igualdade é pedir bug."""
    return int(round(abs(valor) * 100))


def _palavras(norm: str) -> set:
    """Tokens que identificam alguma coisa: só letras, 4+ caracteres, não genéricos."""
    return {p for p in re.findall(r"[a-z]{4,}", norm) if p not in STOPWORDS}


def _parse_valor(texto: str) -> float:
    """
    Aceita `1234.56`, `1.234,56`, `-45,90`, `R$ 45,90` e `(45,90)` (negativo
    contábil). Ambiguidade conhecida: `1.234` sozinho é lido como decimal.
    """
    t = texto.strip().replace("R$", "").replace(" ", "")
    if not t:
        raise ValueError("valor vazio")
    negativo = t.startswith("-") or (t.startswith("(") and t.endswith(")"))
    t = t.strip("()-+")
    if "," in t and "." in t:
        if t.rfind(",") > t.rfind("."):
            t = t.replace(".", "").replace(",", ".")
        else:
            t = t.replace(",", "")
    elif "," in t:
        t = t.replace(",", ".")
    try:
        v = float(t)
    except ValueError:
        raise ValueError("valor inválido: %s" % texto)
    return -v if negativo else v


def _parse_data(texto: str) -> str:
    """
    A mensagem de erro é sempre nossa: o motivo de uma linha inválida chega à
    tela do usuário, e a exceção original do Python (`day 31 must be in range
    1..28`) seria mostrada em inglês.
    """
    t = texto.strip()
    m = re.fullmatch(r"(\d{2})/(\d{2})/(\d{4})", t)
    if m:
        d, mes, a = m.groups()
        return _data_real(a, mes, d, texto)
    m = re.fullmatch(r"(\d{4})-(\d{2})-(\d{2})(?:[ T].*)?", t)
    if m:
        a, mes, d = m.groups()
        return _data_real(a, mes, d, texto)
    raise ValueError("data em formato desconhecido: %s" % texto)


def _data_real(ano: str, mes: str, dia: str, original: str) -> str:
    try:
        return date(int(ano), int(mes), int(dia)).isoformat()
    except ValueError:
        raise ValueError("data inexistente: %s" % original)


def _linha(numero: int, data_iso: str, descricao: str, valor: float, fitid: Optional[str]) -> Dict[str, Any]:
    if valor == 0:
        raise ValueError("valor zero")
    return {
        "linha": numero,
        "data": data_iso,
        "descricao": (descricao or "").strip() or "Sem descrição",
        "valor": valor,
        "type": "income" if valor > 0 else "expense",
        "amount": abs(round(valor, 2)),
        "fitid": fitid or None,
    }


# --------------------------------------------------------------------------- #
# Leitura de arquivo
# --------------------------------------------------------------------------- #


def _decodificar(conteudo: bytes) -> str:
    """UTF-8 primeiro; bancos brasileiros costumam exportar OFX em cp1252."""
    for enc in ("utf-8-sig", "cp1252", "latin-1"):
        try:
            return conteudo.decode(enc)
        except UnicodeDecodeError:
            continue
    raise ArquivoInvalido("Não consegui ler o texto do arquivo.")


def _tag(bloco: str, nome: str) -> Optional[str]:
    m = re.search(r"<%s>\s*([^<\r\n]*)" % nome, bloco, re.I)
    return m.group(1).strip() if m else None


def ler_ofx(texto: str) -> Tuple[List[Dict[str, Any]], List[Dict[str, Any]], Optional[str]]:
    """
    OFX 1.x (SGML, tags sem fechamento) e 2.x (XML) pelo mesmo caminho: o que
    interessa são as tags de dentro de cada STMTTRN, e elas são iguais.
    """
    conta = _tag(texto, "ACCTID")
    blocos = re.findall(
        r"<STMTTRN>(.*?)(?=</STMTTRN>|<STMTTRN>|</BANKTRANLIST>|$)", texto, re.S | re.I
    )
    if not blocos:
        raise ArquivoInvalido("Não encontrei transações neste arquivo OFX.")

    linhas: List[Dict[str, Any]] = []
    invalidas: List[Dict[str, Any]] = []
    for n, b in enumerate(blocos, start=1):
        try:
            posted = _tag(b, "DTPOSTED") or ""
            if not re.match(r"\d{8}", posted):
                raise ValueError("data ausente")
            # Só a data de calendário: o fuso do OFX ("[-3:BRT]") não importa e
            # empurrar o dia por causa dele deslocaria o lançamento.
            data_iso = "%s-%s-%s" % (posted[0:4], posted[4:6], posted[6:8])
            valor = _parse_valor(_tag(b, "TRNAMT") or "")
            partes: List[str] = []
            for campo in ("NAME", "MEMO"):
                v = _tag(b, campo)
                if v and v not in partes:
                    partes.append(v)
            linhas.append(_linha(n, data_iso, " - ".join(partes), valor, _tag(b, "FITID")))
        except ValueError as e:
            invalidas.append({"linha": n, "motivo": str(e)})
    return linhas, invalidas, conta


_ALIASES = {
    "data": {"data", "date", "datalancamento", "dt"},
    "descricao": {"descricao", "historico", "memo", "description", "detalhe"},
    "valor": {"valor", "amount", "value", "quantia"},
    "id": {"id", "fitid", "identificador", "documento"},
}


def _mapear_cabecalho(cabecalho: List[str]) -> Dict[str, int]:
    mapa: Dict[str, int] = {}
    for i, h in enumerate(cabecalho):
        chave = re.sub(r"[^a-z0-9]", "", normalizar(str(h)))
        for campo, apelidos in _ALIASES.items():
            if chave in apelidos and campo not in mapa:
                mapa[campo] = i
    faltando = [c for c in ("data", "descricao", "valor") if c not in mapa]
    if faltando:
        raise ArquivoInvalido(
            "A planilha precisa das colunas data, descricao e valor (faltou: %s)." % ", ".join(faltando)
        )
    return mapa


def _linhas_da_tabela(linhas_brutas: List[List[Any]]) -> Tuple[List[Dict[str, Any]], List[Dict[str, Any]]]:
    if not linhas_brutas:
        raise ArquivoInvalido("O arquivo está vazio.")
    mapa = _mapear_cabecalho([str(c or "") for c in linhas_brutas[0]])

    linhas: List[Dict[str, Any]] = []
    invalidas: List[Dict[str, Any]] = []
    for n, bruta in enumerate(linhas_brutas[1:], start=2):
        if not any(str(c).strip() for c in bruta if c is not None):
            continue
        try:
            cru_data = bruta[mapa["data"]]
            data_iso = (
                cru_data.date().isoformat() if isinstance(cru_data, datetime) else _parse_data(str(cru_data))
            )
            cru_valor = bruta[mapa["valor"]]
            valor = float(cru_valor) if isinstance(cru_valor, (int, float)) else _parse_valor(str(cru_valor))
            fitid = str(bruta[mapa["id"]]).strip() if "id" in mapa and bruta[mapa["id"]] else None
            linhas.append(_linha(n, data_iso, str(bruta[mapa["descricao"]] or ""), valor, fitid))
        except (ValueError, IndexError) as e:
            invalidas.append({"linha": n, "motivo": str(e)})
    return linhas, invalidas


def ler_csv(texto: str) -> Tuple[List[Dict[str, Any]], List[Dict[str, Any]], Optional[str]]:
    amostra = texto[:2000]
    delimitador = ";" if amostra.count(";") >= amostra.count(",") else ","
    brutas = [r for r in csv.reader(io.StringIO(texto), delimiter=delimitador)]
    linhas, invalidas = _linhas_da_tabela(brutas)
    return linhas, invalidas, None


def ler_xlsx(conteudo: bytes) -> Tuple[List[Dict[str, Any]], List[Dict[str, Any]], Optional[str]]:
    try:
        from openpyxl import load_workbook

        wb = load_workbook(io.BytesIO(conteudo), read_only=True, data_only=True)
        brutas = [list(r) for r in wb.active.iter_rows(values_only=True)]
    except Exception as e:  # openpyxl levanta vários tipos para arquivo corrompido
        raise ArquivoInvalido("Não consegui abrir a planilha (%s)." % e.__class__.__name__)
    linhas, invalidas = _linhas_da_tabela(brutas)
    return linhas, invalidas, None


def ler_arquivo(nome: str, conteudo: bytes) -> Tuple[str, List[Dict[str, Any]], List[Dict[str, Any]], Optional[str]]:
    """Devolve (formato, linhas, inválidas, conta_do_arquivo)."""
    ext = (nome.rsplit(".", 1)[-1] if "." in nome else "").lower()
    if ext == "xls":
        raise ArquivoInvalido("Planilha .xls antiga não é suportada: salve como .xlsx ou .csv.")

    if ext == "xlsx":
        return ("xlsx",) + ler_xlsx(conteudo)

    texto = _decodificar(conteudo)
    eh_ofx = ext in ("ofx", "qfx") or "<OFX>" in texto[:3000].upper() or texto.lstrip().startswith("OFXHEADER")
    if eh_ofx:
        return ("ofx",) + ler_ofx(texto)
    if ext in ("csv", "txt") or ext == "":
        return ("csv",) + ler_csv(texto)
    raise ArquivoInvalido("Formato .%s não suportado. Use OFX, CSV ou XLSX." % ext)


# --------------------------------------------------------------------------- #
# Identidade e classificação
# --------------------------------------------------------------------------- #


def _identificar(linhas: List[Dict[str, Any]], conta: Optional[str]) -> None:
    """
    Dá a cada linha um `import_hash` e, quando o banco informa, um `external_id`.

    O hash leva um contador de ocorrência da mesma base dentro do arquivo: duas
    linhas legítimas idênticas no mesmo dia (dois cafés de R$ 8) ficam distintas,
    e reimportar o mesmo arquivo continua gerando os mesmos hashes.
    """
    vistas: Dict[str, int] = {}
    for l in linhas:
        base = "|".join(
            [normalizar(l["descricao"]), str(_centavos(l["valor"])), l["type"], l["data"], conta or ""]
        )
        vistas[base] = vistas.get(base, 0) + 1
        l["import_hash"] = hashlib.sha1(("%s#%d" % (base, vistas[base])).encode("utf-8")).hexdigest()
        if l.get("fitid"):
            l["external_id"] = "%s:%s" % (conta, l["fitid"]) if conta else l["fitid"]
        else:
            l["external_id"] = None


def _descricoes_casam(a: str, b: str) -> bool:
    """
    Igualdade normalizada, ou uma contida na outra, ou palavra-chave em comum.

    A contenção só vale quando a parte menor tem alguma palavra que identifica
    algo — "compra" contida em "compra cartão x" não diz nada.
    """
    if not a or not b:
        return False
    if a == b:
        return True
    menor, maior = (a, b) if len(a) <= len(b) else (b, a)
    if menor in maior and _palavras(menor):
        return True
    return bool(_palavras(a) & _palavras(b))


def _dias_entre(a: str, b: str) -> int:
    da = date.fromisoformat(a)
    db = date.fromisoformat(b)
    return abs((da - db).days)


def classificar(linhas: List[Dict[str, Any]], existentes: List[Dict[str, Any]], destino: Destino) -> None:
    """
    Marca cada linha como `nova`, `ja_importada` ou `suspeita` (com candidato).
    No destino Cartão, a linha positiva fica `tratada_depois`: não é compra e
    não entra em nenhuma regra de duplicidade (estorno e pagamento são de #31/#29).

    Determinístico: sem similaridade por limiar. Um percentual de parecença
    esconderia a regra e produziria decisões que ninguém sabe explicar.
    """
    ext_conhecidos = {t["external_id"] for t in existentes if t.get("external_id")}
    # O hash é identidade só quando o banco NÃO deu uma: FITID é a palavra do
    # banco de que duas linhas são transações distintas, e o hash (que se
    # apoia em descrição, valor e data) não pode contradizê-lo. Duas compras
    # iguais no mesmo dia, cada uma com seu FITID, colidiriam no hash.
    hash_de_quem_nao_tem_fitid = {
        t["import_hash"] for t in existentes if t.get("import_hash") and not t.get("external_id")
    }
    hash_todos = {t["import_hash"] for t in existentes if t.get("import_hash")}

    # Só registro sem FITID pode ser duplicata de uma linha do banco: quem tem
    # FITID já foi identificado pelo próprio banco. E só do mesmo destino.
    candidatos = [
        t
        for t in existentes
        if not t.get("external_id")
        and t.get("ingest_state", "confirmed") == "confirmed"
        and destino.pode_ser_duplicata(t)
    ]
    # Candidato já reivindicado por uma linha que está esperando não é oferecido
    # a outra: o mesmo lançamento não pode ser duplicado por duas linhas.
    usados = {t["reconcile_candidate_id"] for t in existentes if t.get("reconcile_candidate_id")}

    for l in linhas:
        if destino.eh_cartao and l["type"] == "income":
            l["situacao"] = "tratada_depois"
            continue
        if l["external_id"]:
            # Com FITID: ele decide. O hash só vale contra registro que veio
            # sem FITID (ex.: o mesmo extrato importado antes como CSV).
            ja = l["external_id"] in ext_conhecidos or l["import_hash"] in hash_de_quem_nao_tem_fitid
        else:
            ja = l["import_hash"] in hash_todos
        if ja:
            l["situacao"] = "ja_importada"
            continue

        desc = normalizar(l["descricao"])
        melhor: Optional[Dict[str, Any]] = None
        melhor_chave: Optional[Tuple[int, int]] = None
        for t in candidatos:
            if t["id"] in usados or t["type"] != l["type"]:
                continue
            if _centavos(t["amount"]) != _centavos(l["valor"]):
                continue
            distancia = _dias_entre(l["data"], t.get("settled_at") or t["due_date"])
            if distancia > JANELA_DIAS:
                continue
            if not _descricoes_casam(desc, normalizar(t.get("description"))):
                continue
            chave = (distancia, t["id"])
            if melhor_chave is None or chave < melhor_chave:
                melhor, melhor_chave = t, chave

        if melhor is None:
            l["situacao"] = "nova"
        else:
            l["situacao"] = "suspeita"
            l["candidato_id"] = melhor["id"]
            l["candidato"] = {
                "id": melhor["id"],
                "descricao": melhor.get("description"),
                "data": melhor.get("settled_at") or melhor["due_date"],
                "settled": bool(melhor.get("settled_at")),
            }
            usados.add(melhor["id"])


# --------------------------------------------------------------------------- #
# Prévia e importação
# --------------------------------------------------------------------------- #


def analisar(
    user_id: int,
    nome: str,
    conteudo: bytes,
    account_id: Optional[int] = None,
    card_id: Optional[int] = None,
) -> Dict[str, Any]:
    """Lê e classifica sem gravar nada — é o que a prévia mostra."""
    destino = Destino.resolver(user_id, account_id, card_id)
    formato, linhas, invalidas, acctid = ler_arquivo(nome, conteudo)
    if not linhas and not invalidas:
        raise ArquivoInvalido("Não encontrei nenhuma transação neste arquivo.")

    _identificar(linhas, destino.prefixo_hash)
    existentes = [t for t in DataService.load_json("transactions") if t["user_id"] == user_id]
    classificar(linhas, existentes, destino)

    contagem = {"nova": 0, "ja_importada": 0, "suspeita": 0, "tratada_depois": 0}
    entradas = 0.0
    saidas = 0.0
    for l in linhas:
        contagem[l["situacao"]] += 1
        if l["type"] == "income":
            entradas += l["amount"]
        else:
            saidas += l["amount"]

    return {
        "formato": formato,
        "destino": destino.tipo,
        destino.tipo: destino.resumo(),
        "aviso_conta": destino.aviso_de_arquivo_trocado(acctid),
        "arquivo_acctid": acctid,
        "total": len(linhas),
        "contagem": contagem,
        # Totais do arquivo inteiro: é o que denuncia um extrato com o sinal
        # invertido (uma conta cheia de "entradas" que deveriam ser saídas; um
        # cartão cuja soma de compras veio positiva).
        "totais": {"entradas": round(entradas, 2), "saidas": round(saidas, 2)},
        "invalidas": invalidas,
        "linhas": linhas,
    }


ACOES_DECISAO = ("merge", "not_duplicate", "queue")


def _registro_importado(
    l: Dict[str, Any],
    user_id: int,
    destino: Destino,
    novo_id: int,
    agora: str,
    estado: str,
) -> Dict[str, Any]:
    return {
        "id": novo_id,
        "user_id": user_id,
        "type": l["type"],
        "amount": l["amount"],
        "description": l["descricao"],
        # Nulo só nasce de importação: é o "o classificador não soube".
        # A classificação automática é da Fatia 6.
        "category_id": None,
        "due_date": l["data"],
        "settled_at": l["data"],
        "account_id": destino.account_id,
        "card_id": destino.card_id,
        "is_internal_transfer": False,
        "needs_transfer_review": False,
        "series_id": None,
        "series_index": None,
        "ingest_state": estado,
        "source": "import",
        "external_id": l["external_id"],
        "import_hash": l["import_hash"],
        "bank_description": l["descricao"],
        "reconcile_candidate_id": l.get("candidato_id") if estado == "awaiting_reconciliation" else None,
        "history": [
            {"at": agora, "event": "created", "source": "import"},
            {"at": agora, "event": "settled", "on": l["data"]},
        ],
        "created_at": agora,
    }


def _validar_decisoes(analise: Dict[str, Any], decisoes: Optional[Dict[str, Any]]) -> Dict[str, str]:
    """
    Decisão vale só para linha suspeita deste arquivo. Uma decisão que não casa
    com nenhuma é recusada em vez de ignorada: ignorar em silêncio faria o
    usuário achar que decidiu algo que não foi aplicado.
    """
    if not decisoes:
        return {}
    suspeitas = {l["import_hash"] for l in analise["linhas"] if l["situacao"] == "suspeita"}
    limpas: Dict[str, str] = {}
    for chave, acao in decisoes.items():
        if acao not in ACOES_DECISAO:
            raise ArquivoInvalido("Decisão desconhecida: %s." % acao)
        if chave not in suspeitas:
            raise ArquivoInvalido("Há uma decisão para uma linha que não é suspeita de duplicidade neste arquivo.")
        limpas[chave] = acao
    return limpas


@em_transacao
def importar(
    user_id: int,
    nome: str,
    conteudo: bytes,
    account_id: Optional[int] = None,
    decisoes: Optional[Dict[str, Any]] = None,
    card_id: Optional[int] = None,
) -> Dict[str, Any]:
    """
    Grava o que a prévia mostrou.

    Linhas novas entram confirmadas. Para cada linha SUSPEITA o usuário pode ter
    decidido na hora — `merge` funde no candidato, `not_duplicate` a confirma
    como nova — e o que ficar sem decisão (ou com `queue`) espera na fila de
    conciliação. Linhas já importadas são ignoradas.

    Recalcula a partir do arquivo em vez de confiar em linhas devolvidas pelo
    cliente: o que se grava não pode depender de um payload adulterável. A
    decisão é o único dado do cliente aceito, e é validada contra o arquivo.
    """
    destino = Destino.resolver(user_id, account_id, card_id)
    analise = analisar(user_id, nome, conteudo, account_id, card_id)
    escolhidas = _validar_decisoes(analise, decisoes)

    transactions = DataService.load_json("transactions")
    proximo = max((t["id"] for t in transactions), default=0) + 1
    agora = datetime.now(timezone.utc).isoformat().replace("+00:00", "Z")

    novas = 0
    conciliadas = 0
    em_espera = 0
    for l in analise["linhas"]:
        if l["situacao"] in ("ja_importada", "tratada_depois"):
            continue

        if l["situacao"] == "nova":
            transactions.append(_registro_importado(l, user_id, destino, proximo, agora, "confirmed"))
            proximo += 1
            novas += 1
            continue

        acao = escolhidas.get(l["import_hash"], "queue")
        if acao == "merge":
            alvo = next(t for t in transactions if t["id"] == l["candidato_id"])
            registro = _registro_importado(l, user_id, destino, 0, agora, "confirmed")
            _aplicar_fusao(alvo, registro, agora)
            conciliadas += 1
        elif acao == "not_duplicate":
            transactions.append(_registro_importado(l, user_id, destino, proximo, agora, "confirmed"))
            proximo += 1
            novas += 1
        else:
            transactions.append(
                _registro_importado(l, user_id, destino, proximo, agora, "awaiting_reconciliation")
            )
            proximo += 1
            em_espera += 1

    DataService.save_json("transactions", transactions)

    # Primeira importação da conta/cartão: guarda o identificador do banco para
    # poder avisar se um arquivo de outro chegar aqui depois.
    if analise["arquivo_acctid"]:
        registros = DataService.load_json(destino.colecao)
        gravado = next(r for r in registros if r["id"] == destino.id)
        if not gravado.get("file_ref"):
            gravado["file_ref"] = analise["arquivo_acctid"]
            DataService.save_json(destino.colecao, registros)

    return {
        "formato": analise["formato"],
        "destino": destino.tipo,
        destino.tipo: destino.resumo(),
        "total": analise["total"],
        "importadas": novas,
        "conciliadas": conciliadas,
        "ja_existiam": analise["contagem"]["ja_importada"],
        "aguardando_conciliacao": em_espera,
        "tratadas_depois": analise["contagem"]["tratada_depois"],
        "invalidas": analise["invalidas"],
        # Todas as importadas entram sem categoria até a Fatia 6.
        "sem_categoria": novas,
        "classificadas": {"memoria": 0, "ia": 0},
    }


# --------------------------------------------------------------------------- #
# Conciliação
# --------------------------------------------------------------------------- #


def _aplicar_fusao(alvo: Dict[str, Any], linha: Dict[str, Any], agora: str) -> None:
    """
    Funde a linha do banco no registro existente.

    O registro que fica absorve a identidade da linha — é isso que impede a
    mesma linha de voltar como suspeita na próxima reimportação.
    """
    alvo["external_id"] = linha.get("external_id")
    # O hash antigo é MANTIDO quando existe: ele é o que faz o arquivo original
    # (ex.: o CSV) continuar sendo reconhecido. Só herda o da linha se não tinha.
    if not alvo.get("import_hash"):
        alvo["import_hash"] = linha.get("import_hash")
    alvo["bank_description"] = linha.get("bank_description") or linha.get("description")
    if alvo.get("account_id") is None:
        alvo["account_id"] = linha.get("account_id")

    evento: Dict[str, Any] = {
        "at": agora,
        "event": "reconciled_with",
        "transaction_id": linha.get("id"),
        "bank_description": alvo["bank_description"],
    }
    if _centavos(alvo["amount"]) != _centavos(linha["amount"]):
        # Escolha manual com valor diferente (juros, tarifa): o valor do usuário
        # é mantido, mas a diferença fica registrada em vez de sumir.
        evento["amount_bank"] = linha["amount"]
    alvo.setdefault("history", []).append(evento)

    # Previsto vira realizado: o banco confirma que o dinheiro se moveu. Se o
    # usuário já tinha efetivado, a data dele prevalece — é edição dele.
    if not alvo.get("settled_at"):
        alvo["settled_at"] = linha.get("settled_at") or linha["due_date"]
        alvo["history"].append({"at": agora, "event": "settled", "on": alvo["settled_at"]})


@em_transacao
def conciliar(transaction_id: int, action: str, with_id: Optional[int] = None) -> Dict[str, Any]:
    """
    Resolve uma linha em espera.

    `merge` funde na direção do registro existente; `not_duplicate` a confirma
    como transação nova. Levanta `LookupError` quando algo não existe e
    `ValueError` quando o pedido não faz sentido.
    """
    transactions = DataService.load_json("transactions")
    linha = next((t for t in transactions if t["id"] == transaction_id), None)
    if linha is None:
        raise LookupError("Lançamento não encontrado.")
    if linha.get("ingest_state") != "awaiting_reconciliation":
        raise ValueError("Este lançamento não está aguardando conciliação.")

    agora = datetime.now(timezone.utc).isoformat().replace("+00:00", "Z")

    if action == "not_duplicate":
        linha["ingest_state"] = "confirmed"
        linha["reconcile_candidate_id"] = None
        linha.setdefault("history", []).append({"at": agora, "event": "confirmed_not_duplicate"})
        DataService.save_json("transactions", transactions)
        return linha

    if action != "merge":
        raise ValueError("Ação desconhecida: %s" % action)

    alvo_id = with_id or linha.get("reconcile_candidate_id")
    if not alvo_id:
        raise ValueError("Informe com qual lançamento conciliar.")
    alvo = next((t for t in transactions if t["id"] == alvo_id), None)
    if alvo is None:
        raise LookupError("O lançamento escolhido para conciliar não existe mais.")
    if alvo["id"] == linha["id"]:
        raise ValueError("Não dá para conciliar um lançamento com ele mesmo.")
    if alvo.get("external_id"):
        # Absorver sobrescreveria o FITID que o banco já deu a ele.
        raise ValueError("Este lançamento já tem identificador do banco; não há o que absorver.")
    if alvo.get("account_id") not in (None, linha.get("account_id")):
        raise ValueError("Este lançamento pertence a outra conta.")
    if alvo.get("card_id") != linha.get("card_id"):
        raise ValueError("Este lançamento pertence a outro destino (conta ou cartão).")

    _aplicar_fusao(alvo, linha, agora)
    transactions = [t for t in transactions if t["id"] != linha["id"]]
    DataService.save_json("transactions", transactions)
    return alvo


# --------------------------------------------------------------------------- #
# Modelo padrão para download
# --------------------------------------------------------------------------- #

# O layout da aplicação: o que o sistema precisa, e nada além. Sem etapa de
# mapeamento de colunas — quem tem um arquivo de outro formato o ajusta a este.
COLUNAS_MODELO = ("data", "descricao", "valor")

INSTRUCOES_MODELO = [
    ("data", "Dia do lançamento, em DD/MM/AAAA (ex.: 05/09/2026) ou AAAA-MM-DD."),
    ("descricao", "Texto do lançamento, como aparece no extrato."),
    ("valor", "Com sinal: NEGATIVO é saída, POSITIVO é entrada. Vírgula ou ponto decimal (ex.: -152,30)."),
]

EXEMPLOS_MODELO = [
    ("05/09/2026", "Supermercado Extra", "-152,30"),
    ("10/09/2026", "Salário setembro", "3500,00"),
]


def modelo_padrao(formato: str) -> Tuple[bytes, str, str]:
    """
    (conteúdo, tipo MIME, nome do arquivo) do modelo.

    A aba/planilha de dados sai **só com o cabeçalho**: exemplos ali seriam
    importados por engano por quem enviasse o modelo sem apagá-los. Os exemplos
    ficam na aba de instruções do XLSX.
    """
    if formato == "csv":
        # UTF-8 com BOM e ponto e vírgula: é o que o Excel em português abre
        # sem quebrar acento nem juntar tudo numa coluna só.
        texto = ";".join(COLUNAS_MODELO) + "\r\n"
        return texto.encode("utf-8-sig"), "text/csv; charset=utf-8", "modelo-importacao.csv"

    if formato == "xlsx":
        from openpyxl import Workbook

        wb = Workbook()
        dados = wb.active
        dados.title = "Lançamentos"
        dados.append(list(COLUNAS_MODELO))
        for col, largura in zip("ABC", (14, 44, 14)):
            dados.column_dimensions[col].width = largura

        ajuda = wb.create_sheet("Como preencher")
        ajuda.append(["Coluna", "O que colocar"])
        for coluna, texto in INSTRUCOES_MODELO:
            ajuda.append([coluna, texto])
        ajuda.append([])
        ajuda.append(["Exemplos (não são importados — ficam só nesta aba)"])
        ajuda.append(list(COLUNAS_MODELO))
        for linha in EXEMPLOS_MODELO:
            ajuda.append(list(linha))
        ajuda.column_dimensions["A"].width = 16
        ajuda.column_dimensions["B"].width = 80

        wb.active = 0  # a importação lê a aba ativa: precisa ser a de dados
        buf = io.BytesIO()
        wb.save(buf)
        return (
            buf.getvalue(),
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            "modelo-importacao.xlsx",
        )

    raise ValueError("Formato de modelo desconhecido: %s" % formato)
