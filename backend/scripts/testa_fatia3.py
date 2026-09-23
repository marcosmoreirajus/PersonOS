# -*- coding: utf-8 -*-
"""
Critérios de aceite da Fatia 3 (docs/finance/PRD.md), verificados sobre uma
cópia dos dados — o arquivo real não é tocado.

    importar um arquivo com 3 linhas repetidas de um extrato já importado
    resulta em 0 duplicatas gravadas;
    conciliar uma linha com um lançamento manual e reimportar o mesmo arquivo
    não gera nova suspeita.

Além dos dois critérios do PRD, cobre as decisões de implementação que o spec
deixava em aberto (ver o cabeçalho de import_service.py).
"""

import shutil
import sys
import tempfile
from datetime import date, timedelta
from pathlib import Path

RAIZ = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(RAIZ))

from app.services import data_service  # noqa: E402
from app.services import import_service as imp  # noqa: E402
from app.services.data_service import DataService  # noqa: E402

falhas = []


def checa(nome, condicao, detalhe=""):
    print(("  OK    " if condicao else "  FALHA ") + nome + (" -- " + detalhe if detalhe else ""))
    if not condicao:
        falhas.append(nome)


def tx():
    return DataService.load_json("transactions")


def visiveis():
    """O que saldo, relatório e listagem padrão enxergam."""
    return DataService.get_transactions_by_user(1)


def ofx_sgml(linhas, conta="12345-6"):
    """OFX 1.x de banco brasileiro: SGML sem fechamento, cp1252."""
    corpo = "".join(
        "<STMTTRN>\r\n<TRNTYPE>DEBIT\r\n<DTPOSTED>%s000000[-3:BRT]\r\n<TRNAMT>%s\r\n"
        "<FITID>%s\r\n<MEMO>%s\r\n</STMTTRN>\r\n" % (d.replace("-", ""), v, fitid, memo)
        for (d, v, fitid, memo) in linhas
    )
    texto = (
        "OFXHEADER:100\r\nDATA:OFXSGML\r\nVERSION:102\r\nCHARSET:1252\r\n\r\n"
        "<OFX><BANKMSGSRSV1><STMTTRNRS><STMTRS><BANKACCTFROM><ACCTID>%s</ACCTID></BANKACCTFROM>"
        "<BANKTRANLIST>%s</BANKTRANLIST></STMTRS></STMTTRNRS></BANKMSGSRSV1></OFX>" % (conta, corpo)
    )
    return texto.encode("cp1252")


def csv_bytes(linhas, sep=";"):
    cab = sep.join(["data", "descricao", "valor"])
    corpo = "\n".join(sep.join(l) for l in linhas)
    return (cab + "\n" + corpo + "\n").encode("utf-8")


def main():
    tmp = Path(tempfile.mkdtemp(prefix="clari-fatia3-"))
    for f in (RAIZ / "data").glob("*.json"):
        shutil.copy(f, tmp / f.name)
    data_service.DATA_DIR = tmp
    print("dados de teste em", tmp)
    print()

    hoje = date.today()
    d = lambda n: (hoje - timedelta(days=n)).isoformat()  # noqa: E731

    # A importação exige uma conta: quem chega escolhe de onde o extrato veio.
    conta_id = DataService.create_account(1, "Bradesco", "checking")["id"]
    analisar = lambda nome, conteudo, conta=None: imp.analisar(1, nome, conteudo, conta or conta_id)  # noqa: E731
    importar = lambda nome, conteudo, conta=None, decisoes=None: imp.importar(  # noqa: E731
        1, nome, conteudo, conta or conta_id, decisoes)

    # ------------------------------------------------------------ leitura
    print("Leitura")
    formato, linhas, invalidas, conta = imp.ler_arquivo(
        "extrato.ofx",
        ofx_sgml(
            [
                (d(10), "-152,30", "A001", "COMPRA SUPERMERCADO EXTRA"),
                (d(9), "3500.00", "A002", "SALÁRIO SETEMBRO"),
                (d(8), "-8.00", "A003", "PADARIA PÃO QUENTE"),
            ]
        ),
    )
    checa("OFX SGML em cp1252 é lido", formato == "ofx" and len(linhas) == 3, "%d linhas" % len(linhas))
    checa("conta do arquivo vem do ACCTID", conta == "12345-6")
    checa("acento cp1252 preservado", linhas[1]["descricao"] == "SALÁRIO SETEMBRO", linhas[1]["descricao"])
    checa("vírgula e ponto decimais são aceitos", linhas[0]["amount"] == 152.30 and linhas[1]["amount"] == 3500.0)
    checa("sinal define o tipo", linhas[0]["type"] == "expense" and linhas[1]["type"] == "income")
    checa("fuso do OFX não desloca o dia", linhas[0]["data"] == d(10), linhas[0]["data"])

    _, csv_linhas, _, _ = imp.ler_arquivo(
        "extrato.csv",
        csv_bytes([["10/09/2026", "Mercado", "-1.234,56"], ["2026-09-11", "Freela", "500,00"]]),
    )
    checa("CSV com data BR, ISO e valor 1.234,56", csv_linhas[0]["valor"] == -1234.56 and csv_linhas[1]["data"] == "2026-09-11")
    _, csv_v, _, _ = imp.ler_arquivo("e.csv", csv_bytes([["10/09/2026", "Mercado", "-10.50"]], sep=","))
    checa("CSV separado por vírgula", csv_v[0]["valor"] == -10.5)

    try:
        imp.ler_arquivo("a.csv", b"foo;bar\n1;2\n")
        checa("planilha sem as colunas é recusada", False)
    except imp.ArquivoInvalido as e:
        checa("planilha sem as colunas é recusada com mensagem clara", "data" in str(e), str(e))

    _, _, invs, _ = imp.ler_arquivo(
        "e.csv", csv_bytes([["10/09/2026", "Ok", "-5,00"], ["31/02/2026", "Data ruim", "-1,00"]])
    )
    checa("linha inválida é reportada, não engolida", len(invs) == 1 and invs[0]["linha"] == 3, str(invs))

    # --------------------------------------------- prévia não grava nada
    print()
    print("Importação")
    arquivo_a = ofx_sgml([(d(20 - i), "-%d.00" % (10 + i), "F%03d" % i, "COMPRA LOJA %s" % chr(65 + i)) for i in range(10)])
    antes = len(tx())
    prev = analisar("a.ofx", arquivo_a)
    checa("prévia não grava", len(tx()) == antes)
    checa("prévia conta as linhas", prev["contagem"]["nova"] == 10, str(prev["contagem"]))

    # ----------------------------------------------------- critério do PRD 1
    r1 = importar("a.ofx", arquivo_a)
    checa("importa as 10 linhas", r1["importadas"] == 10 and r1["aguardando_conciliacao"] == 0, str(r1))
    importadas = [t for t in tx() if t.get("source") == "import"]
    checa("linha importada nasce efetivada e sem categoria",
          all(t["settled_at"] == t["due_date"] and t["category_id"] is None for t in importadas))
    checa("linha importada carrega identidade de origem",
          all(t["external_id"] and t["import_hash"] for t in importadas))

    arquivo_b = ofx_sgml(
        [(d(20 - i), "-%d.00" % (10 + i), "F%03d" % i, "COMPRA LOJA %s" % chr(65 + i)) for i in range(3)]
        + [(d(2), "-77.00", "N001", "FARMÁCIA NOVA"), (d(1), "-88.00", "N002", "LIVRARIA NOVA")]
    )
    antes = len(tx())
    r2 = importar("b.ofx", arquivo_b)
    checa("PRD: 3 linhas repetidas de extrato já importado -> 0 duplicatas",
          r2["ja_existiam"] == 3 and r2["importadas"] == 2, str(r2))
    checa("só as 2 novas foram gravadas", len(tx()) - antes == 2, "+%d" % (len(tx()) - antes))

    antes = len(tx())
    r3 = importar("a.ofx", arquivo_a)
    checa("reimportar o mesmo arquivo não grava nada", r3["importadas"] == 0 and len(tx()) == antes, str(r3))

    # ------------------------------------ linhas idênticas legítimas (CSV)
    cafes = csv_bytes([[d(5), "Cafe da esquina", "-8,00"], [d(5), "Cafe da esquina", "-8,00"]])
    r4 = importar("cafes.csv", cafes)
    checa("dois cafés iguais no mesmo dia entram os dois", r4["importadas"] == 2, str(r4))
    r5 = importar("cafes.csv", cafes)
    checa("e reimportar continua idempotente", r5["importadas"] == 0 and r5["ja_existiam"] == 2, str(r5))

    # ------------------------------------------------------- suspeita
    print()
    print("Duplicidade")
    manual = DataService.create_transaction(
        user_id=1, category_id=1, type="expense", amount=152.30,
        due_date=d(11), settled_at=d(11), description="Supermercado Extra",
    )
    extrato = ofx_sgml([(d(10), "-152,30", "S001", "COMPRA SUPERMERCADO EXTRA"),
                        (d(3), "-19,90", "S002", "ASSINATURA STREAMING")])
    prev = analisar("s.ofx", extrato)
    checa("lançamento manual vira candidato (valor igual, ±3 dias, palavra em comum)",
          prev["linhas"][0]["situacao"] == "suspeita" and prev["linhas"][0]["candidato_id"] == manual["id"],
          str(prev["contagem"]))

    saldo_antes = round(sum(t["amount"] for t in visiveis() if t["type"] == "expense" and t.get("settled_at")), 2)
    r6 = importar("s.ofx", extrato)
    checa("entrada parcial: a limpa entra na hora, a suspeita espera",
          r6["importadas"] == 1 and r6["aguardando_conciliacao"] == 1, str(r6))
    em_espera = [t for t in tx() if t.get("ingest_state") == "awaiting_reconciliation"]
    checa("a suspeita fica na mesma tabela, em espera", len(em_espera) == 1)
    checa("em espera NÃO aparece em saldo nem listagem padrão",
          all(t["id"] != em_espera[0]["id"] for t in visiveis()))
    saldo_depois = round(sum(t["amount"] for t in visiveis() if t["type"] == "expense" and t.get("settled_at")), 2)
    checa("saldo só mudou pela linha limpa (19,90)", round(saldo_depois - saldo_antes, 2) == 19.90,
          "%.2f -> %.2f" % (saldo_antes, saldo_depois))

    r7 = importar("s.ofx", extrato)
    checa("reimportar com a suspeita esperando não duplica a fila",
          r7["importadas"] == 0 and r7["aguardando_conciliacao"] == 0 and r7["ja_existiam"] == 2, str(r7))

    # ----------------------------------------------- critério do PRD 2
    linha = em_espera[0]
    ok = imp.conciliar(linha["id"], "merge")
    checa("fundir devolve o registro que fica", ok["id"] == manual["id"])
    checa("o que fica mantém id e categoria", ok["category_id"] == 1)
    checa("o que fica absorve os identificadores do banco",
          ok["external_id"] and ok["import_hash"] and ok["bank_description"] == "COMPRA SUPERMERCADO EXTRA")
    checa("a linha em espera some", all(t["id"] != linha["id"] for t in tx()))
    checa("evento reconciled_with entra no histórico",
          any(e["event"] == "reconciled_with" for e in ok["history"]))
    r8 = importar("s.ofx", extrato)
    checa("PRD: conciliar e reimportar o mesmo arquivo não gera nova suspeita",
          r8["aguardando_conciliacao"] == 0 and r8["importadas"] == 0 and r8["ja_existiam"] == 2, str(r8))

    # ----------------------------------------------- previsto vira realizado
    print()
    print("Previsto vira realizado")
    serie_tx = DataService.create_transaction(
        user_id=1, category_id=5, type="expense", amount=1800.0,
        due_date=d(2), description="Aluguel",
        series={"kind": "recurring", "frequency": "monthly"},
    )
    prevista = next(t for t in tx() if t.get("series_id") == serie_tx["series_id"] and t["due_date"] == d(2))
    checa("ocorrência de série nasce prevista (sem efetivação)", prevista["settled_at"] is None)
    importar("al.ofx", ofx_sgml([(d(1), "-1800.00", "AL1", "ALUGUEL SETEMBRO")]))
    espera = [t for t in tx() if t.get("ingest_state") == "awaiting_reconciliation"]
    checa("linha do extrato casa com a ocorrência PREVISTA da série",
          len(espera) == 1 and espera[0]["reconcile_candidate_id"] == prevista["id"])
    fund = imp.conciliar(espera[0]["id"], "merge")
    checa("fundir efetiva a ocorrência: previsto vira realizado",
          fund["settled_at"] == d(1) and fund["series_id"] == serie_tx["series_id"], str(fund["settled_at"]))

    # ------------------------- candidato já efetivado mantém a data do usuário
    manual2 = DataService.create_transaction(
        user_id=1, category_id=2, type="expense", amount=40.0,
        due_date=d(6), settled_at=d(6), description="Uber centro",
    )
    importar("ub.ofx", ofx_sgml([(d(4), "-40.00", "UB1", "UBER CENTRO")]))
    esp2 = [t for t in tx() if t.get("ingest_state") == "awaiting_reconciliation"]
    f2 = imp.conciliar(esp2[0]["id"], "merge")
    checa("registro já efetivado mantém a data escolhida pelo usuário", f2["settled_at"] == d(6), f2["settled_at"])

    # -------------------------------------------------------- não é duplicata
    manual3 = DataService.create_transaction(
        user_id=1, category_id=3, type="expense", amount=60.0,
        due_date=d(7), settled_at=d(7), description="Farmacia Sao Joao",
    )
    importar("fa.ofx", ofx_sgml([(d(7), "-60.00", "FA1", "FARMACIA SAO JOAO")]))
    esp3 = [t for t in tx() if t.get("ingest_state") == "awaiting_reconciliation"]
    antes_saldo = len(visiveis())
    conf = imp.conciliar(esp3[0]["id"], "not_duplicate")
    checa("'não é duplicata' confirma e a linha entra normalmente",
          conf["ingest_state"] == "confirmed" and len(visiveis()) == antes_saldo + 1)

    # ----------------------------------------------- casos que NÃO devem suspeitar
    print()
    print("Falso positivo")
    DataService.create_transaction(
        user_id=1, category_id=1, type="expense", amount=33.0,
        due_date=d(12), settled_at=d(12), description="Compra cartao",
    )
    p = analisar("g.ofx", ofx_sgml([(d(12), "-33.00", "G1", "COMPRA CARTAO PADARIA")]))
    checa("palavra genérica ('compra', 'cartão') não basta pra suspeitar", p["linhas"][0]["situacao"] == "nova")

    # duas compras iguais que vieram do banco (FITIDs diferentes) são distintas
    importar("d1.ofx", ofx_sgml([(d(13), "-25.00", "D1", "RESTAURANTE SABOR")]))
    p = analisar("d2.ofx", ofx_sgml([(d(13), "-25.00", "D2", "RESTAURANTE SABOR")]))
    checa("registro que já veio do banco não é candidato a duplicata",
          p["linhas"][0]["situacao"] == "nova", p["linhas"][0]["situacao"])

    # cada candidato é reivindicado por uma linha só
    DataService.create_transaction(
        user_id=1, category_id=1, type="expense", amount=50.0,
        due_date=d(14), settled_at=d(14), description="Padaria do Ze",
    )
    p = analisar("u.ofx", ofx_sgml([(d(14), "-50.00", "U1", "PADARIA DO ZE"),
                                          (d(14), "-50.00", "U2", "PADARIA DO ZE")]))
    sit = sorted(l["situacao"] for l in p["linhas"])
    checa("duas linhas iguais, um lançamento manual: uma suspeita e uma nova", sit == ["nova", "suspeita"], str(sit))

    # tipo e valor diferentes não casam
    DataService.create_transaction(
        user_id=1, category_id=1, type="expense", amount=99.0,
        due_date=d(15), settled_at=d(15), description="Livraria Cultura",
    )
    p = analisar("v.ofx", ofx_sgml([(d(15), "99.00", "V1", "LIVRARIA CULTURA")]))
    checa("mesmo valor mas tipo oposto (crédito x débito) não casa", p["linhas"][0]["situacao"] == "nova")
    p = analisar("w.ofx", ofx_sgml([(d(15), "-99.01", "W1", "LIVRARIA CULTURA")]))
    checa("diferença de 1 centavo não casa", p["linhas"][0]["situacao"] == "nova")
    p = analisar("x.ofx", ofx_sgml([(d(20), "-99.00", "X1", "LIVRARIA CULTURA")]))
    checa("fora da janela de ±3 dias não casa", p["linhas"][0]["situacao"] == "nova")

    # ------------------------------------------- entre formatos e planilha
    print()
    print("Formatos")
    mesmas = [(d(30), "-45.90", "CF1", "ACADEMIA BODY TECH"), (d(29), "-12.00", "CF2", "ESTACIONAMENTO CENTRO")]
    importar("cf.csv", csv_bytes([[l[0], l[3], l[1].replace(".", ",")] for l in mesmas]))
    r_of = importar("cf.ofx", ofx_sgml(mesmas))
    checa("mesmo extrato em CSV e depois OFX, na MESMA conta, é reconhecido como já importado",
          r_of["importadas"] == 0 and r_of["aguardando_conciliacao"] == 0 and r_of["ja_existiam"] == 2, str(r_of))

    outra_conta = DataService.create_account(1, "C6 Bank", "checking")["id"]
    r_c6 = importar("cf.ofx", ofx_sgml(mesmas), conta=outra_conta)
    checa("a mesma linha em OUTRA conta é distinta (uma transferência aparece nas duas pontas)",
          r_c6["importadas"] == 2, str(r_c6))
    checa("cada lançamento importado carrega a conta escolhida",
          {t["account_id"] for t in tx() if t.get("bank_description") in ("ACADEMIA BODY TECH",)} == {conta_id, outra_conta})

    from io import BytesIO
    from openpyxl import Workbook
    wb = Workbook()
    ws = wb.active
    ws.append(["Data", "Histórico", "Valor"])
    ws.append([date.fromisoformat(d(40)), "Pagamento fornecedor", -300.5])
    ws.append([date.fromisoformat(d(39)), "Recebimento cliente", 1200])
    buf = BytesIO()
    wb.save(buf)
    formato_x, linhas_x, _, _ = imp.ler_arquivo("planilha.xlsx", buf.getvalue())
    checa("xlsx com datas reais e cabeçalho 'Histórico' é lido",
          formato_x == "xlsx" and len(linhas_x) == 2 and linhas_x[0]["amount"] == 300.5, str(len(linhas_x)))
    checa("xlsx: tipo pelo sinal", linhas_x[0]["type"] == "expense" and linhas_x[1]["type"] == "income")

    # ---------------------------------------------------------- contas
    print()
    print("Contas")
    try:
        DataService.create_account(1, "  bradesco ", "checking")
        checa("nome de conta repetido (sem caixa nem espaços) é recusado", False)
    except ValueError:
        checa("nome de conta repetido (sem caixa nem espaços) é recusado", True)
    conta_btg = DataService.create_account(1, "BTG", "checking")
    renomeada = DataService.update_account(conta_btg["id"], name="BTG Pactual")
    checa("renomear conta", renomeada["name"] == "BTG Pactual")
    LOGO_OK = "data:image/png;base64,iVBORw0KGgo="
    c_extra = DataService.create_account(1, "Poupança Teste", "savings", 1234.56, LOGO_OK)
    checa("conta guarda tipo poupança, saldo e logo",
          c_extra["kind"] == "savings" and c_extra["initial_balance"] == 1234.56 and c_extra["logo"] == LOGO_OK, str(c_extra)[:90])
    try:
        DataService.create_account(1, "Logo ruim", "checking", 0, "data:text/html;base64,AAAA")
        checa("logo que não é imagem", False)
    except ValueError:
        checa("logo que não é imagem é recusado", True)
    try:
        DataService.create_account(1, "Logo grande", "checking", 0, "data:image/png;base64," + "A" * 400000)
        checa("logo grande", False)
    except ValueError:
        checa("logo grande demais é recusado", True)
    sem_logo = DataService.update_account(c_extra["id"], logo="")
    checa("string vazia remove o logo e o saldo é mantido", sem_logo["logo"] is None and sem_logo["initial_balance"] == 1234.56)
    checa("atualizar só o saldo", DataService.update_account(c_extra["id"], initial_balance=10)["initial_balance"] == 10.0)
    try:
        DataService.update_account(conta_btg["id"], name="Bradesco")
        checa("renomear para nome em uso é recusado", False)
    except ValueError:
        checa("renomear para nome em uso é recusado", True)
    try:
        imp.analisar(1, "x.csv", csv_bytes([[d(3), "X", "-1,00"]]), 99999)
        checa("importar para conta inexistente", False)
    except LookupError:
        checa("importar para conta inexistente levanta LookupError", True)
    de_outro = DataService.create_account(2, "Conta de outro usuário", "checking")["id"]
    try:
        imp.analisar(1, "x.csv", csv_bytes([[d(3), "X", "-1,00"]]), de_outro)
        checa("conta de outro usuário", False)
    except LookupError:
        checa("conta de OUTRO usuário não pode receber a importação", True)

    # --------------------------------------- aviso de arquivo de outra conta
    conta_acct = DataService.create_account(1, "Conta com identificador", "checking")["id"]
    importar("ac1.ofx", ofx_sgml([(d(70), "-5.00", "AC1", "TARIFA UM")], conta="111-1"), conta=conta_acct)
    p_ok = analisar("ac2.ofx", ofx_sgml([(d(69), "-6.00", "AC2", "TARIFA DOIS")], conta="111-1"), conta=conta_acct)
    p_troca = analisar("ac3.ofx", ofx_sgml([(d(68), "-7.00", "AC3", "TARIFA TRES")], conta="999-9"), conta=conta_acct)
    checa("arquivo com o mesmo identificador da conta não avisa", p_ok["aviso_conta"] is None)
    checa("arquivo de OUTRO banco na mesma conta avisa antes de importar",
          bool(p_troca["aviso_conta"]) and "outra conta" in p_troca["aviso_conta"], str(p_troca["aviso_conta"]))
    p_csv = analisar("ac4.csv", csv_bytes([[d(60), "Sem identificador", "-1,00"]]), conta=conta_acct)
    checa("CSV não tem identificador, então não há o que comparar", p_csv["aviso_conta"] is None)

    # ------------------------------------------------------------ totais
    p_tot = analisar("t.csv", csv_bytes([[d(80), "Salario", "1000,00"], [d(79), "Mercado", "-250,50"], [d(78), "Luz", "-100,00"]]))
    checa("prévia traz os totais de entradas e saídas (denuncia sinal invertido)",
          p_tot["totais"] == {"entradas": 1000.0, "saidas": 350.5}, str(p_tot["totais"]))

    # ------------------------------------------- decisões no commit
    print()
    print("Decisões na importação")
    ms = [
        DataService.create_transaction(user_id=1, category_id=1, type="expense", amount=v,
                                       due_date=d(90), settled_at=d(90), description=desc)
        for (v, desc) in [(71.0, "Oficina Mecanica Silva"), (72.0, "Papelaria Central Norte"), (73.0, "Clinica Vet Amigo")]
    ]
    ext_dec = ofx_sgml([(d(90), "-71.00", "DC1", "OFICINA MECANICA SILVA"),
                        (d(90), "-72.00", "DC2", "PAPELARIA CENTRAL NORTE"),
                        (d(90), "-73.00", "DC3", "CLINICA VET AMIGO")])
    pv = analisar("dc.ofx", ext_dec)
    hs = {l["descricao"]: l["import_hash"] for l in pv["linhas"]}
    checa("três linhas suspeitas", pv["contagem"]["suspeita"] == 3, str(pv["contagem"]))

    antes_invalidas = len(tx())
    try:
        importar("dc.ofx", ext_dec, decisoes={"hash-que-nao-existe": "merge"})
        checa("decisão para linha que não é suspeita é recusada", False)
    except imp.ArquivoInvalido:
        checa("decisão para linha que não é suspeita é recusada (não some em silêncio)", True)
    try:
        importar("dc.ofx", ext_dec, decisoes={hs["OFICINA MECANICA SILVA"]: "explodir"})
        checa("ação de decisão desconhecida", False)
    except imp.ArquivoInvalido:
        checa("ação de decisão desconhecida é recusada", True)
    checa("decisões inválidas não gravaram nada", len(tx()) == antes_invalidas, "%d -> %d" % (antes_invalidas, len(tx())))

    antes = len(tx())
    r_dec = importar("dc.ofx", ext_dec, decisoes={
        hs["OFICINA MECANICA SILVA"]: "merge",
        hs["PAPELARIA CENTRAL NORTE"]: "not_duplicate",
        # a terceira fica sem decisão: vai para a fila
    })
    checa("decidido na importação: 1 fundida, 1 nova, 1 na fila",
          r_dec["conciliadas"] == 1 and r_dec["importadas"] == 1 and r_dec["aguardando_conciliacao"] == 1, str(r_dec))
    fundido = next(t for t in tx() if t["id"] == ms[0]["id"])
    checa("merge aplicado na hora: o existente absorveu o FITID e nenhuma linha extra foi gravada",
          str(fundido["external_id"]).endswith(":DC1") and fundido["bank_description"] == "OFICINA MECANICA SILVA"
          and len(tx()) - antes == 2, "+%d" % (len(tx()) - antes))
    checa("o lançamento fundido ganhou a conta da importação", fundido["account_id"] == conta_id)
    checa("o que ficou sem decisão está na fila, e só ele",
          [t["description"] for t in tx() if t.get("ingest_state") == "awaiting_reconciliation"
           and t["source"] == "import" and "CLINICA" in t["description"]] == ["CLINICA VET AMIGO"])
    r_again = importar("dc.ofx", ext_dec)
    checa("reimportar depois das decisões não gera nada",
          r_again["importadas"] == 0 and r_again["conciliadas"] == 0 and r_again["aguardando_conciliacao"] == 0
          and r_again["ja_existiam"] == 3, str(r_again))

    # ---------------------------------------------------------- modelo
    print()
    print("Modelo padrão")
    csv_modelo, mime_csv, nome_csv = imp.modelo_padrao("csv")
    checa("modelo CSV: UTF-8 com BOM e só o cabeçalho",
          csv_modelo.startswith(b"\xef\xbb\xbf") and csv_modelo.decode("utf-8-sig").strip() == "data;descricao;valor")
    try:
        analisar("modelo-importacao.csv", csv_modelo)
        checa("modelo vazio", False)
    except imp.ArquivoInvalido as e:
        checa("importar o modelo sem preencher dá mensagem clara em vez de importar exemplos",
              "nenhuma transação" in str(e), str(e))

    from openpyxl import load_workbook
    from io import BytesIO as _BIO

    xlsx_modelo, _, nome_xlsx = imp.modelo_padrao("xlsx")
    wb_m = load_workbook(_BIO(xlsx_modelo))
    checa("modelo XLSX: aba de dados só com cabeçalho e instruções em outra aba",
          wb_m.sheetnames == ["Lançamentos", "Como preencher"] and wb_m["Lançamentos"].max_row == 1, str(wb_m.sheetnames))
    wb_m["Lançamentos"].append([date.fromisoformat(d(95)), "Teste do modelo", -42.5])
    buf_m = _BIO()
    wb_m.save(buf_m)
    _, l_modelo, _, _ = imp.ler_arquivo("modelo.xlsx", buf_m.getvalue())
    checa("modelo preenchido é lido, e os exemplos da aba de instruções NÃO são importados",
          len(l_modelo) == 1 and l_modelo[0]["descricao"] == "Teste do modelo", str(len(l_modelo)))

    # ------------------------------------------------------- erros
    print()
    print("Erros")
    try:
        imp.conciliar(999999, "merge")
        checa("conciliar id inexistente", False)
    except LookupError:
        checa("conciliar id inexistente levanta LookupError", True)
    try:
        imp.conciliar(manual["id"], "merge")
        checa("conciliar lançamento que não está em espera", False)
    except ValueError:
        checa("conciliar lançamento que não está em espera é recusado", True)
    try:
        imp.ler_arquivo("a.xls", b"x")
        checa(".xls recusado", False)
    except imp.ArquivoInvalido as e:
        checa(".xls antigo é recusado com orientação", "xlsx" in str(e))

    # ------------------------------------------------------------- API
    # Multipart, códigos HTTP e mensagens: o que o serviço sozinho não exercita.
    print()
    print("API")
    from fastapi.testclient import TestClient
    from app.main import app

    cli = TestClient(app)
    api_ofx = ofx_sgml([(d(50), "-64.90", "API1", "LOJA DA API UM"), (d(49), "-11.10", "API2", "LOJA DA API DOIS")])

    r = cli.post("/api/import/preview", data={"user_id": 1, "account_id": conta_id}, files={"file": ("api.ofx", api_ofx)})
    checa("POST /api/import/preview devolve 200 e a contagem",
          r.status_code == 200 and r.json()["data"]["contagem"]["nova"] == 2, str(r.status_code))
    antes = len(tx())
    checa("a prévia via API não grava", len(tx()) == antes)

    r = cli.post("/api/import/commit", data={"user_id": 1, "account_id": conta_id}, files={"file": ("api.ofx", api_ofx)})
    checa("POST /api/import/commit grava", r.status_code == 200 and r.json()["data"]["importadas"] == 2)
    r = cli.post("/api/import/commit", data={"user_id": 1, "account_id": conta_id}, files={"file": ("api.ofx", api_ofx)})
    checa("reimportar via API é idempotente", r.json()["data"]["importadas"] == 0)

    r = cli.post("/api/import/preview", data={"user_id": 1, "account_id": conta_id}, files={"file": ("x.pdf", b"%PDF-1.4 nada")})
    checa("formato não suportado devolve 400 com mensagem em português",
          r.status_code == 400 and "Formato" in r.json()["detail"], r.text[:80])
    r = cli.post("/api/import/preview", data={"user_id": 1, "account_id": conta_id}, files={"file": ("vazio.csv", b"")})
    checa("arquivo vazio devolve 400", r.status_code == 400)
    r = cli.post("/api/import/preview", data={"user_id": 1, "account_id": conta_id},
                 files={"file": ("grande.csv", b"data;descricao;valor\n" + b"x" * (5 * 1024 * 1024 + 10))})
    checa("arquivo acima de 5 MB devolve 413", r.status_code == 413, str(r.status_code))

    # lista a fila via filtro e resolve via /api/reconcile
    manual_api = DataService.create_transaction(
        user_id=1, category_id=1, type="expense", amount=30.0,
        due_date=d(60), settled_at=d(60), description="Mercadinho Central",
    )
    cli.post("/api/import/commit", data={"user_id": 1, "account_id": conta_id},
             files={"file": ("m.ofx", ofx_sgml([(d(60), "-30.00", "M1", "MERCADINHO CENTRAL")]))})
    fila_toda = cli.get("/api/transactions/user/1", params={"ingest_state": "awaiting_reconciliation"}).json()["data"]
    # A fila pode ter itens de outros casos deste script: o caso isola o que é dele.
    fila = [x for x in fila_toda if x["reconcile_candidate_id"] == manual_api["id"]]
    checa("GET ?ingest_state=awaiting_reconciliation devolve a fila",
          len(fila) == 1 and all(x["ingest_state"] == "awaiting_reconciliation" for x in fila_toda))
    padrao = cli.get("/api/transactions/user/1").json()["data"]
    checa("a listagem padrão NÃO traz a fila", all(x.get("ingest_state", "confirmed") == "confirmed" for x in padrao))

    r = cli.post("/api/reconcile/%d" % fila[0]["id"], json={"action": "merge"})
    checa("POST /api/reconcile funde e devolve o registro que fica",
          r.status_code == 200 and r.json()["data"]["id"] == manual_api["id"])
    r = cli.post("/api/reconcile/%d" % fila[0]["id"], json={"action": "merge"})
    checa("conciliar de novo a mesma linha devolve 404 (ela já foi absorvida)", r.status_code == 404, str(r.status_code))
    r = cli.post("/api/reconcile/%d" % manual_api["id"], json={"action": "merge"})
    checa("conciliar quem não está em espera devolve 400", r.status_code == 400, str(r.status_code))
    r = cli.post("/api/reconcile/1", json={"action": "explodir"})
    checa("ação desconhecida é rejeitada na validação (422)", r.status_code == 422, str(r.status_code))

    # modelo para download
    r = cli.get("/api/import/template/csv")
    checa("GET /api/import/template/csv devolve o arquivo como anexo",
          r.status_code == 200 and "attachment" in r.headers.get("content-disposition", "")
          and r.content.startswith(b"\xef\xbb\xbf"), str(r.status_code))
    r = cli.get("/api/import/template/xlsx")
    checa("GET /api/import/template/xlsx devolve um xlsx válido", r.status_code == 200 and r.content[:2] == b"PK")
    r = cli.get("/api/import/template/pdf")
    checa("modelo em formato desconhecido devolve 404", r.status_code == 404)

    # contas
    r = cli.post("/api/accounts", json={"user_id": 1, "name": "Conta API", "kind": "card"})
    conta_api = r.json()["data"] if r.status_code == 200 else {}
    checa("POST /api/accounts cria a conta", r.status_code == 200 and conta_api.get("kind") == "card", r.text[:80])
    r = cli.post("/api/accounts", json={"user_id": 1, "name": "conta  api", "kind": "card"})
    checa("nome repetido devolve 400 com mensagem", r.status_code == 400 and "nome" in r.json()["detail"], r.text[:80])
    r = cli.post("/api/accounts", json={"user_id": 1, "name": "X", "kind": "poupanca"})
    checa("tipo de conta inválido é rejeitado na validação (422)", r.status_code == 422, str(r.status_code))
    lista = cli.get("/api/accounts/user/1").json()["data"]
    checa("GET /api/accounts/user lista as contas em ordem alfabética",
          [a["name"] for a in lista] == sorted([a["name"] for a in lista], key=str.lower) and len(lista) >= 4)
    r = cli.patch("/api/accounts/%d" % conta_api["id"], json={"name": "Conta API renomeada"})
    checa("PATCH /api/accounts renomeia", r.status_code == 200 and r.json()["data"]["name"] == "Conta API renomeada")
    r = cli.patch("/api/accounts/99999", json={"name": "Nada"})
    checa("PATCH em conta inexistente devolve 404", r.status_code == 404)

    # importação: conta e decisões pela API
    r = cli.post("/api/import/preview", data={"user_id": 1, "account_id": 99999}, files={"file": ("q.ofx", api_ofx)})
    checa("prévia com conta inexistente devolve 404", r.status_code == 404, str(r.status_code))
    r = cli.post("/api/import/commit", data={"user_id": 1, "account_id": conta_id, "decisoes": "{nao json"},
                 files={"file": ("q.ofx", api_ofx)})
    checa("decisões em JSON inválido devolvem 400", r.status_code == 400, str(r.status_code))
    r = cli.post("/api/import/commit", data={"user_id": 1, "account_id": conta_id, "decisoes": "[]"},
                 files={"file": ("q.ofx", api_ofx)})
    checa("decisões que não são um objeto devolvem 400", r.status_code == 400, str(r.status_code))
    r = cli.post("/api/import/commit", data={"user_id": 1, "account_id": conta_id, "decisoes": "{}"},
                 files={"file": ("q.ofx", api_ofx)})
    checa("commit com decisões vazias funciona (tudo já importado)", r.status_code == 200
          and r.json()["data"]["importadas"] == 0)

    print()
    if falhas:
        print("FALHARAM: %d" % len(falhas))
        for f_ in falhas:
            print("  -", f_)
        return 1
    print("todos os critérios passaram")
    shutil.rmtree(tmp, ignore_errors=True)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
